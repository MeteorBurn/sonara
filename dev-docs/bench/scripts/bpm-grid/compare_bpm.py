"""Compare Sonara, Beat This! and Mixed In Key BPM (and Sonara's beat grid).

Reads (all read-only):
  data/broken_1000.sqlite        Sonara (bpm, beats, downbeats)
  data/beat_this/NNNN - *.json   Beat This! (bpm, beats_sec, downbeats_sec)
  data/mik/NNNN - *.json         Mixed In Key (bpm)
Writes data/bpm_comparison.xlsx and data/bpm_comparison.json.

Reference = MIK BPM, on tracks where Beat This! (basic 60/median-interval BPM)
is within AGREE_REL of MIK, i.e. both chose the same metrical level.
Categories of a value against a reference:
  exact <=0.05 | fine <=0.5 | minor <=2 | x2 x0.5 x1.5 x2/3 x4/3 x3/4 | major
  critical = any multiple or major (> 2 BPM).
Grid vs Beat This! (first 5 s ignored, as mir_eval): beat F-measure +-70 ms,
median signed phase offset (Sonara - Beat This!) in ms, downbeat F-measure.
"""

from __future__ import annotations

import json
import sqlite3
import sys
import zlib
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from beat_this_bpm import bpm_estimates, fold_to_range  # noqa: E402
import bpm_grid_paths as P  # noqa: E402

DATA = P.DATA
AGREE_REL = 0.03
TOL = 0.07
SKIP = 5.0
MULTIPLES = [(2.0, "x2"), (0.5, "x0.5"), (1.5, "x1.5"), (2 / 3, "x2/3"),
             (4 / 3, "x4/3"), (0.75, "x3/4")]
ORDER = ["exact", "fine", "minor", "x2", "x0.5", "x1.5", "x2/3", "x4/3", "x3/4",
         "major", "n/a"]


def category(value, ref):
    if value is None or ref is None:
        return "n/a"
    d = abs(value - ref)
    if d <= 0.05:
        return "exact"
    if d <= 0.5:
        return "fine"
    if d <= 2.0:
        return "minor"
    ratio = value / ref
    for m, label in MULTIPLES:
        if abs(ratio / m - 1.0) <= 0.02:
            return label
    return "major"


def is_critical(cat):
    return cat not in ("exact", "fine", "minor", "n/a")


def f_measure(est, ref):
    est = est[est >= SKIP]
    ref = ref[ref >= SKIP]
    if len(est) == 0 or len(ref) == 0:
        return None, None
    used = np.zeros(len(est), dtype=bool)
    offsets = []
    for r in ref:
        i = int(np.searchsorted(est, r))
        best = None
        for j in (i - 1, i):
            if 0 <= j < len(est) and not used[j] and abs(est[j] - r) <= TOL:
                if best is None or abs(est[j] - r) < abs(est[best] - r):
                    best = j
        if best is not None:
            used[best] = True
            offsets.append(est[best] - r)
    hits = len(offsets)
    p, rc = hits / len(est), hits / len(ref)
    f = 0.0 if hits == 0 else 2 * p * rc / (p + rc)
    return f, (float(np.median(offsets)) * 1000.0 if offsets else None)


def load_json_dir(folder: Path) -> dict[int, dict]:
    out = {}
    for f in folder.glob("[0-9][0-9][0-9][0-9] - *.json"):
        d = json.loads(f.read_text(encoding="utf-8"))
        out[int(d["idx"])] = d
    return out


def sonara_rows(db: Path) -> dict[int, dict]:
    con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
    out = {}
    for path, pidx, status, error, sj in con.execute(
            "SELECT path, playlist_idx, status, error, scalars_json FROM tracks"):
        row = {"path": path, "status": status, "error": error}
        if sj:
            s = json.loads(sj)
            prov = s.get("provenance", {})
            hop, sr = prov.get("hop_length", 512), prov.get("sample_rate", 22050)
            row.update(bpm=s.get("bpm"), bpm_raw=s.get("bpm_raw"),
                       bpm_conf=s.get("bpm_confidence"),
                       reg=s.get("rhythmic_regularity"),
                       reg_conf=s.get("rhythmic_regularity_confidence"),
                       grid_stability=s.get("grid_stability"),
                       duration=s.get("duration_sec"))
            arrays = {n: np.frombuffer(zlib.decompress(d), dtype=dt).reshape(json.loads(sh))
                      for n, dt, sh, d in con.execute(
                          "SELECT name, dtype, shape, data FROM arrays WHERE path = ? "
                          "AND name IN ('beats', 'downbeats')", (path,))}
            for n in ("beats", "downbeats"):
                frames = arrays.get(n)
                if frames is None and isinstance(s.get(n), list):
                    frames = np.asarray(s[n])
                row[n] = None if frames is None else frames.astype(float) * hop / sr
        out[pidx + 1] = row
    return out


def main() -> int:
    son = sonara_rows(DATA / "broken_1000.sqlite")
    bt = load_json_dir(DATA / "beat_this")
    mik = load_json_dir(DATA / "mik")
    n = max(len(bt), len(mik), len(son))
    rows = []
    for idx in range(1, n + 1):
        s, b, m = son.get(idx, {}), bt.get(idx, {}), mik.get(idx, {})
        path = b.get("path") or m.get("path") or s.get("path")
        s_bpm, b_bpm, m_bpm = s.get("bpm"), b.get("bpm"), m.get("bpm")
        ibi = b.get("ibi_median_sec")
        b_val = fold_to_range(60.0 / ibi, 79.0, 192.0) if ibi else b_bpm
        m_val = m.get("tempo_raw") if m.get("bpm") is not None else None
        bm_cat = category(b_val, m_val)
        agree = (b_val is not None and m_val is not None
                 and abs(b_val / m_val - 1.0) <= AGREE_REL)
        ref = m_val if agree else None
        s_cat = category(s_bpm, ref) if agree else "n/a"
        s_vs_bt, s_vs_mik = category(s_bpm, b_val), category(s_bpm, m_val)
        # Sonara's own beats, same basic rule as Beat This! (60 / median interval)
        s_beats_fit = None
        if s.get("beats") is not None and len(s["beats"]) >= 3:
            med = bpm_estimates(np.asarray(s["beats"]))["bpm_median"]
            s_beats_fit = None if med is None else fold_to_range(med, 79.0, 192.0)
        beat_f = phase = down_f = None
        if s.get("beats") is not None and b.get("beats_sec"):
            beat_f, phase = f_measure(np.asarray(s["beats"]), np.asarray(b["beats_sec"]))
        if s.get("downbeats") is not None and b.get("downbeats_sec"):
            down_f, _ = f_measure(np.asarray(s["downbeats"]), np.asarray(b["downbeats_sec"]))
        rows.append({
            "idx": idx, "file": Path(path).name if path else None, "path": path,
            "sonara_status": s.get("status"), "sonara_bpm": s_bpm,
            "sonara_beats_fit_bpm": s_beats_fit, "sonara_bpm_conf": s.get("bpm_conf"),
            "beat_this_bpm": b_val, "beat_this_raw": b.get("bpm_raw"),
            "mik_bpm": m_val, "bt_minus_mik": None if bm_cat == "n/a" else b_val - m_val,
            "refs": "agree" if agree else ("n/a" if bm_cat == "n/a" else "disagree"),
            "bt_vs_mik": bm_cat, "ref_bpm": ref,
            "sonara_minus_ref": None if ref is None or s_bpm is None else s_bpm - ref,
            "sonara_cat": s_cat, "sonara_critical": is_critical(s_cat),
            "sonara_vs_bt": s_vs_bt, "sonara_vs_mik": s_vs_mik,
            "beat_f": beat_f, "phase_ms": phase, "downbeat_f": down_f,
            "grid_stability": s.get("grid_stability"),
            "regularity": s.get("reg"), "regularity_conf": s.get("reg_conf"),
            "duration_sec": s.get("duration"), "sonara_error": s.get("error"),
        })

    # ---------------- statistics ----------------
    agree_rows = [r for r in rows if r["refs"] == "agree"]
    bm = np.abs([r["bt_minus_mik"] for r in agree_rows])
    spread = {f"p{q}": float(np.percentile(bm, q)) for q in (50, 90, 95, 99)}
    spread["max"] = float(bm.max())

    def counts(key, subset):
        c = {k: 0 for k in ORDER}
        for r in subset:
            c[r[key]] += 1
        return {k: v for k, v in c.items() if v}

    s_ok = [r for r in agree_rows if r["sonara_bpm"] is not None]
    stats = {
        "tracks": len(rows),
        "with_bpm": {"sonara": sum(r["sonara_bpm"] is not None for r in rows),
                     "beat_this": sum(r["beat_this_bpm"] is not None for r in rows),
                     "mik": sum(r["mik_bpm"] is not None for r in rows)},
        "refs": {k: sum(r["refs"] == k for r in rows) for k in ("agree", "disagree", "n/a")},
        "bt_vs_mik": counts("bt_vs_mik", rows),
        "bt_mik_abs_diff_on_agree": spread,
        "sonara_vs_ref": counts("sonara_cat", agree_rows),
        "sonara_vs_bt_all": counts("sonara_vs_bt", rows),
        "sonara_vs_mik_all": counts("sonara_vs_mik", rows),
        "sonara_critical_on_ref": sum(r["sonara_critical"] for r in agree_rows),
        "sonara_abs_err_on_ref_exactfine": {
            f"p{q}": float(np.percentile(np.abs([r["sonara_minus_ref"] for r in s_ok
                                                 if r["sonara_cat"] in ("exact", "fine")]), q))
            for q in (50, 90, 95, 99)},
    }
    good_tempo = [r for r in agree_rows if r["sonara_cat"] in ("exact", "fine")
                  and r["beat_f"] is not None]
    if good_tempo:
        bf = np.array([r["beat_f"] for r in good_tempo])
        df = np.array([r["downbeat_f"] for r in good_tempo if r["downbeat_f"] is not None])
        ph = np.array([r["phase_ms"] for r in good_tempo if r["phase_ms"] is not None])
        stats["grid_on_correct_tempo"] = {
            "tracks": len(good_tempo),
            "beat_f_median": float(np.median(bf)),
            "beat_f_ge_0.9": int((bf >= 0.9).sum()), "beat_f_lt_0.5": int((bf < 0.5).sum()),
            "downbeat_f_median": float(np.median(df)) if len(df) else None,
            "downbeat_f_ge_0.9": int((df >= 0.9).sum()), "downbeat_f_lt_0.5": int((df < 0.5).sum()),
            "phase_ms_median": float(np.median(ph)), "phase_ms_p10": float(np.percentile(ph, 10)),
            "phase_ms_p90": float(np.percentile(ph, 90)),
        }
    (DATA / "bpm_comparison.json").write_text(
        json.dumps({"stats": stats, "rows": rows}, ensure_ascii=False, indent=1, default=float),
        encoding="utf-8")

    # ---------------- workbook ----------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    bold = Font(bold=True)
    lines = [
        ("Sonara vs Beat This! vs Mixed In Key — BPM and beat grid, 1000 broken-rhythm tracks", None),
        ("Reference: MIK BPM where Beat This! is within ±3% of MIK (same metrical level); grid reference: Beat This! beats", None),
        ("Categories: exact <=0.05 | fine <=0.5 | minor <=2 | multiples (x2, x0.5, x1.5, x2/3, x4/3, x3/4) | major >2; critical = multiple or major", None),
        ("", None),
    ]
    for text, _ in lines:
        ws.append([text])
    ws["A1"].font = Font(bold=True, size=13)

    def block(title, data: dict, fmt=None):
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        for k, v in data.items():
            ws.append([f"  {k}", v])
            if fmt and isinstance(v, float):
                ws.cell(ws.max_row, 2).number_format = fmt
        ws.append([])

    block("Tracks with BPM", stats["with_bpm"])
    block("Reference status (Beat This! vs MIK)", stats["refs"])
    block("Beat This! vs MIK categories (all tracks)", stats["bt_vs_mik"])
    block("|Beat This! - MIK| on agreeing tracks, BPM (Beat This! basic BPM is frame-quantized)", spread, "0.000")
    block("Sonara vs reference (agreeing tracks)", stats["sonara_vs_ref"])
    block("|Sonara - reference| where exact/fine, BPM", stats["sonara_abs_err_on_ref_exactfine"], "0.000")
    block("Sonara vs Beat This! (all tracks)", stats["sonara_vs_bt_all"])
    block("Sonara vs MIK (all tracks)", stats["sonara_vs_mik_all"])
    if "grid_on_correct_tempo" in stats:
        block("Sonara grid vs Beat This! (tracks where Sonara tempo is exact/fine)",
              stats["grid_on_correct_tempo"], "0.00")
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14

    cols = [
        ("idx", "#", 6, "0"), ("file", "File", 48, None),
        ("sonara_bpm", "Sonara BPM", 11, "0.00"),
        ("beat_this_bpm", "Beat This! BPM", 13, "0.00"),
        ("mik_bpm", "MIK BPM", 10, "0.00"),
        ("ref_bpm", "Reference BPM", 13, "0.00"),
        ("sonara_minus_ref", "Sonara - ref", 12, "+0.00;-0.00;0.00"),
        ("sonara_cat", "Sonara vs ref", 13, None),
        ("sonara_critical", "Critical", 9, None),
        ("refs", "Refs", 10, None),
        ("bt_minus_mik", "BT - MIK", 10, "+0.00;-0.00;0.00"),
        ("bt_vs_mik", "BT vs MIK", 11, None),
        ("sonara_vs_bt", "Sonara vs BT", 12, None),
        ("sonara_vs_mik", "Sonara vs MIK", 13, None),
        ("sonara_beats_fit_bpm", "Sonara beats median BPM", 13, "0.00"),
        ("sonara_bpm_conf", "Sonara BPM conf", 11, "0.00"),
        ("beat_f", "Beat F (±70ms)", 11, "0.00"),
        ("phase_ms", "Phase ms", 10, "+0;-0;0"),
        ("downbeat_f", "Downbeat F", 11, "0.00"),
        ("grid_stability", "Grid stability", 11, "0.00"),
        ("regularity", "Regularity", 10, "0.00"),
        ("regularity_conf", "Regularity conf", 11, "0.00"),
        ("duration_sec", "Duration s", 10, "0"),
        ("beat_this_raw", "BT raw BPM", 10, "0.00"),
        ("sonara_status", "Sonara status", 11, None),
        ("sonara_error", "Sonara error", 30, None),
        ("path", "Path", 80, None),
    ]
    red = PatternFill("solid", fgColor="F4C7C3")
    amber = PatternFill("solid", fgColor="FCE8B2")
    green = PatternFill("solid", fgColor="D9EAD3")

    def sheet(title, subset, sort_key=None):
        w = wb.create_sheet(title)
        w.append([c[1] for c in cols])
        for cell in w[1]:
            cell.font = bold
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for r in (sorted(subset, key=sort_key) if sort_key else subset):
            w.append([r[c[0]] if not isinstance(r[c[0]], np.generic) else float(r[c[0]])
                      for c in cols])
        for i, (_, _, width, fmt) in enumerate(cols, 1):
            letter = get_column_letter(i)
            w.column_dimensions[letter].width = width
            if fmt:
                for cell in w[letter][1:]:
                    cell.number_format = fmt
        last = w.max_row
        if last > 1:
            for key in ("sonara_cat", "bt_vs_mik", "sonara_vs_bt", "sonara_vs_mik"):
                L = get_column_letter([c[0] for c in cols].index(key) + 1)
                rng = f"{L}2:{L}{last}"
                w.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"exact"'], fill=green))
                w.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'OR({L}2="fine",{L}2="minor")'], fill=amber))
                w.conditional_formatting.add(rng, FormulaRule(
                    formula=[f'AND({L}2<>"exact",{L}2<>"fine",{L}2<>"minor",{L}2<>"n/a",{L}2<>"")'],
                    fill=red))
            for key in ("beat_f", "downbeat_f"):
                L = get_column_letter([c[0] for c in cols].index(key) + 1)
                rng = f"{L}2:{L}{last}"
                w.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["0.5"], fill=red))
                w.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.5", "0.9"], fill=amber))
        w.freeze_panes = "C2"
        w.auto_filter.ref = w.dimensions
        return w

    sheet("All tracks", rows)
    sheet("Sonara critical", [r for r in rows if r["sonara_critical"]],
          lambda r: (r["sonara_cat"], r["idx"]))
    sheet("Refs disagree", [r for r in rows if r["refs"] == "disagree"],
          lambda r: (r["bt_vs_mik"], r["idx"]))
    sheet("Grid issues", [r for r in good_tempo if (r["beat_f"] or 0) < 0.9
                          or (r["downbeat_f"] is not None and r["downbeat_f"] < 0.5)],
          lambda r: (r["beat_f"] or 0))
    sheet("Not exact", [r for r in agree_rows if r["sonara_cat"] in ("fine", "minor")],
          lambda r: -abs(r["sonara_minus_ref"] or 0))
    out = DATA / "bpm_comparison.xlsx"
    wb.save(out)
    print(json.dumps(stats, indent=1, ensure_ascii=False))
    print(f"\nwritten: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
