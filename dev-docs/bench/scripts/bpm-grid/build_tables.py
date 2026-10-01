"""Build the per-program bench tables from the per-track JSON.

For every program folder in data/ (1_sonara, 2_mixed_in_key, 3_rekordbox,
4_beat_this) and set (broken, straight) this writes <prog>-<set>.sqlite (tables
run, fields, tracks, arrays) and <prog>-<set>.xlsx (per-track values, formatted).
The input is json_<set>/ only, which is never modified. One exception: the JSON
does not carry Sonara's onset_strength_bands, so they are copied from the
existing sonara-<set>.sqlite.

BPM fields are named the same for every program:
    bpm           the program's BPM rounded to 2 decimals
    bpm_raw       the BPM as the program reports it
    bpm_unfolded  (Sonara, Beat This!) the tempo before folding into 79-192

Only values a program stores, or exact transforms of them, are written.

    python build_tables.py [--only sonara|mik|rekordbox|beat_this ...] [--out-dir DIR]
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import os
import re
import sqlite3
import sys
import zlib
from decimal import Decimal
from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import bpm_grid_paths as P  # noqa: E402
from beat_this_bpm import beats_per_bar, bpm_estimates, fold_to_range  # noqa: E402

SETS = ("broken", "straight")
BPM_LO, BPM_HI = 79.0, 192.0
AMBER = PatternFill("solid", fgColor="FCE8B2")
GREEN = PatternFill("solid", fgColor="D9EAD3")
BAND = PatternFill("solid", fgColor="EEF3F8")
WHITE_LINE = Side(style="thin", color="FFFFFF")


def now_utc() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def file_name(path: str) -> str:
    return re.split(r"[\\/]", path.strip())[-1]


def track_files(folder: Path) -> list[tuple[int, dict]]:
    """(idx, record) per `0001 - <stem>.json`; idx is the 1-based file prefix."""
    out = []
    for f in sorted(folder.glob("[0-9][0-9][0-9][0-9] - *.json")):
        out.append((int(f.name[:4]), json.loads(f.read_text(encoding="utf-8"))))
    return out


def read_meta(folder: Path) -> dict:
    meta = folder / "_run_meta.json"
    return json.loads(meta.read_text(encoding="utf-8")) if meta.is_file() else {}


def encode(values, dtype: str) -> tuple[str, str, bytes]:
    arr = np.asarray(values, dtype=dtype)
    return arr.dtype.str, json.dumps(list(arr.shape)), zlib.compress(arr.tobytes())


def r2(x):
    return None if x is None else round(x, 2)


# ---------------------------------------------------------------------------
# Sonara
# ---------------------------------------------------------------------------
SONARA_COPY = ("bpm_confidence", "n_beats", "tempo_variability", "grid_offset_sec", "grid_stability",
               "time_signature", "time_signature_confidence", "meter_grid_offset_sec",
               "meter_grid_stability", "rhythmic_regularity", "rhythmic_regularity_label",
               "rhythmic_regularity_confidence", "meter_rhythmic_regularity",
               "meter_rhythmic_regularity_label", "meter_rhythmic_regularity_confidence",
               "duration_sec")
SONARA_ARRAYS = {"beats": "<i4", "downbeats": "<i4", "meter_downbeats": "<i4",
                 "tempo_curve": "<f4", "onset_frames": "<i4"}
SONARA_RUN_KEYS = ("sonara_version", "sonara_repo_head", "wheel_sha256", "playlist",
                   "playlist_sha256", "playlist_tracks", "last_started_at")


def load_sonara(set_name: str, folder: Path):
    jdir = folder / f"json_{set_name}"
    meta = read_meta(jdir)
    bands_db = folder / f"sonara-{set_name}.sqlite"
    if not bands_db.is_file():
        sys.exit(f"{bands_db} is missing: it is the only copy of onset_strength_bands")
    con = sqlite3.connect(f"file:{bands_db}?mode=ro&immutable=1", uri=True)
    bands = {i: (d, s, b) for i, d, s, b in con.execute(
        "SELECT idx, dtype, shape, data FROM arrays WHERE name = 'onset_strength_bands'")}
    con.close()
    run, rows, arrays = None, [], {}
    for idx, t in track_files(jdir):
        prov = t.get("provenance") or {}
        if run is None:
            run = {"set": set_name, "program": "Sonara", **{k: meta.get(k) for k in SONARA_RUN_KEYS},
                   "schema_version": prov.get("schema_version"), "mode": prov.get("mode"),
                   "sample_rate": prov.get("sample_rate"), "hop_length": prov.get("hop_length"),
                   "bpm_min": prov.get("bpm_min"), "bpm_max": prov.get("bpm_max"),
                   "onset_band_edges_hz": t.get("onset_band_edges_hz"),
                   "frame_to_sec": "sec = frame * hop_length / sample_rate"}
        row = {"idx": idx, "file_name": file_name(t["path"]), "path": t["path"], "status": t.get("status"),
               "bpm": r2(t.get("bpm")), "bpm_raw": t.get("bpm"), "bpm_unfolded": t.get("bpm_raw"),
               **{k: t.get(k) for k in SONARA_COPY}}
        cands = t.get("bpm_candidates") or []
        for i in range(5):
            row[f"cand{i + 1}_bpm"], row[f"cand{i + 1}_score"] = cands[i] if i < len(cands) else (None, None)
        for k in ("rhythmic_regularity_candidates", "meter_rhythmic_regularity_candidates"):
            row[k] = None if t.get(k) is None else json.dumps(t[k], ensure_ascii=False)
        rows.append(row)
        arrays[idx] = {n: encode(t[n], d) for n, d in SONARA_ARRAYS.items() if isinstance(t.get(n), list)}
        if idx not in bands:
            sys.exit(f"no onset_strength_bands for {set_name} idx {idx} in {bands_db}")
        arrays[idx]["onset_strength_bands"] = bands[idx]
    run["source"] = f"{jdir.name}/ + onset_strength_bands from the previous {bands_db.name}"
    return run, rows, arrays


SONARA = {
    "prefix": "sonara", "dir": P.SONARA_DIR, "load": load_sonara,
    "groups": {"Track": "1F4E78", "BPM": "2E75B6", "Tempo candidates": "5B9BD5", "Beat grid": "548235",
               "Meter": "7F6000", "Rhythmic regularity": "7030A0", "File": "595959"},
    "columns": [
        ("idx", "Track", "bench playlist position (1-1000)", "0"),
        ("file_name", "Track", "file name", None),
        ("bpm", "BPM", "Sonara bpm rounded to 2 decimals", "0.00"),
        ("bpm_raw", "BPM", "Sonara bpm as reported (folded into bpm_min-bpm_max, 79-192)", "0.00000"),
        ("bpm_unfolded", "BPM", "Sonara's own bpm_raw: the tempo before range folding", "0.00000"),
        ("bpm_confidence", "BPM", "bpm_confidence (developer formula)", "0.000"),
        *[c for i in range(1, 6) for c in (
            (f"cand{i}_bpm", "Tempo candidates", f"bpm_candidates[{i}]: candidate BPM (Sonara order)", "0.000"),
            (f"cand{i}_score", "Tempo candidates", f"bpm_candidates[{i}]: candidate score", "0.000"))],
        ("n_beats", "Beat grid", "number of beats", "0"),
        ("tempo_variability", "Beat grid", "tempo variability across beats", "0.0000"),
        ("grid_offset_sec", "Beat grid", "beat grid start, s", "0.000"),
        ("grid_stability", "Beat grid", "beat grid stability (0-1)", "0.000"),
        ("time_signature", "Meter", "time signature", None),
        ("time_signature_confidence", "Meter", "time signature confidence", "0.000"),
        ("meter_grid_offset_sec", "Meter", "meter grid start, s", "0.000"),
        ("meter_grid_stability", "Meter", "meter grid stability (0-1)", "0.000"),
        ("rhythmic_regularity", "Rhythmic regularity", "regularity score; empty = measured, undecidable", "0.000"),
        ("rhythmic_regularity_label", "Rhythmic regularity", "regular / irregular", None),
        ("rhythmic_regularity_confidence", "Rhythmic regularity", "score confidence", "0.000"),
        ("meter_rhythmic_regularity", "Rhythmic regularity", "same on the meter grid", "0.000"),
        ("meter_rhythmic_regularity_label", "Rhythmic regularity", "regular / irregular", None),
        ("meter_rhythmic_regularity_confidence", "Rhythmic regularity", "score confidence", "0.000"),
        ("duration_sec", "File", "duration, s", "0.00"),
        ("status", "File", "analysis status", None),
        ("path", "File", "file path", None),
    ],
    "sqlite_only": [
        ("rhythmic_regularity_candidates", "Rhythmic regularity", "[[label, probability], ...] (JSON)"),
        ("meter_rhythmic_regularity_candidates", "Rhythmic regularity", "same on the meter grid (JSON)"),
    ],
    "arrays": {
        "beats": ("frame", "beat frames"),
        "downbeats": ("frame", "downbeat frames"),
        "meter_downbeats": ("frame", "meter downbeat frames"),
        "tempo_curve": ("bpm", "tempo between adjacent beats (length n_beats - 1)"),
        "onset_strength_bands": ("per frame", "onset envelope: 5 bands x frames (edges in run.onset_band_edges_hz)"),
        "onset_frames": ("frame", "onset frames"),
    },
    "notes": [("(amber bpm_unfolded)", "", "bpm_unfolded differs from bpm_raw: Sonara folded the tempo", "xlsx")],
}


# ---------------------------------------------------------------------------
# Mixed In Key
# ---------------------------------------------------------------------------
def timespan_sec(value: str) -> float:
    """.NET TimeSpan 'hh:mm:ss.fffffff' -> seconds, exact to the stored 100 ns tick."""
    h, m, s = value.split(":")
    return float(Decimal(h) * 3600 + Decimal(m) * 60 + Decimal(s))


def load_mik(set_name: str, folder: Path):
    jdir = folder / f"json_{set_name}"
    meta = read_meta(jdir)
    rows, arrays = [], {}
    for idx, r in track_files(jdir):
        st = r.get("structure") or {}
        seg = [(timespan_sec(x["StartTime"]), timespan_sec(x["EndTime"])) for x in st.get("EnergySegments") or []]
        aes = st.get("AnalysisEnergySegments") or []
        tempo = r.get("tempo_raw")
        rows.append({
            "idx": idx, "file_name": file_name(r["path"]), "path": r["path"],
            "bpm": r2(tempo), "bpm_raw": tempo, "grid_start_sec": seg[0][0] if seg else None,
            "duration_sec": timespan_sec(aes[-1]["EndTime"]) if aes else None,
            "status": r.get("status"), "analyzed_at": r.get("last_analyzed_utc"),
            "mik_song_id": r.get("mik_song_id"), "file_path_hash": r.get("file_path_hash")})
        if seg:
            arrays[idx] = {"segments_sec": encode(seg, "<f8")}
    times = sorted(r["analyzed_at"] for r in rows if r["analyzed_at"])
    run = {"set": set_name, "program": "Mixed In Key 11",
           "bpm_range": "Restrict range 79-192 (MIK settings, from the user; not stored by MIK)",
           "bpm_precision": "Song.Tempo is a double; MIK shows 2 decimals",
           "analyzed_utc_range": f"{times[0]} .. {times[-1]}" if times else None,
           "mik_db": meta.get("source"), "playlist": meta.get("playlist"),
           "volume_serials": meta.get("volume_serials"), "exported_at": meta.get("exported_at"),
           "source": f"{jdir.name}/ (mik_export.py)"}
    return run, rows, arrays


MIK = {
    "prefix": "mik", "dir": P.MIK_DIR, "load": load_mik,
    "groups": {"Track": "1F4E78", "BPM": "2E75B6", "Grid": "548235", "File": "595959", "MIK": "7F6000"},
    "columns": [
        ("idx", "Track", "bench playlist position (1-1000)", "0"),
        ("file_name", "Track", "file name", None),
        ("bpm", "BPM", "Song.Tempo rounded to 2 decimals, as MIK displays it", "0.00"),
        ("bpm_raw", "BPM", "Song.Tempo as MIK stores it (double)", "0.000000"),
        ("grid_start_sec", "Grid", "start of the first EnergySegment = anchor of MIK's beat grid "
                                   "(every segment boundary sits on this grid)", "0.0000"),
        ("duration_sec", "File", "end of the last AnalysisEnergySegment (analysed length)", "0.000"),
        ("status", "File", "ok = analysed by MIK (IsAnalyzed and Tempo > 0)", None),
        ("analyzed_at", "File", "Song.LastAnalyzedUtc", None),
        ("mik_song_id", "MIK", "Song.Id", None),
        ("file_path_hash", "MIK", "Song.FilePathHash (MIK's key for the file)", None),
        ("path", "File", "Song.File, the path MIK analysed", None),
    ],
    "sqlite_only": [],
    "arrays": {"segments_sec": ("sec", "EnergySegments [start, end] in seconds; boundaries lie on MIK's beat "
                                       "grid, most on bar lines")},
    "notes": [],
}


# ---------------------------------------------------------------------------
# Rekordbox
# ---------------------------------------------------------------------------
def load_rekordbox(set_name: str, folder: Path):
    jdir = folder / f"json_{set_name}"
    meta = read_meta(jdir)
    rows, arrays = [], {}
    for idx, t in track_files(jdir):
        attrs = t.get("attributes") or {}
        marks = t.get("tempo") or []
        raw = float(attrs["AverageBpm"]) if attrs.get("AverageBpm") else None
        bpms = [m["Bpm"] for m in marks]
        rows.append({
            "idx": idx, "file_name": file_name(t["path"]), "path": t["path"],
            "bpm": r2(raw) if raw else None, "bpm_raw": raw,
            "tempo_min": min(bpms) if bpms else None, "tempo_max": max(bpms) if bpms else None,
            "n_tempo_markers": len(marks), "grid_start_sec": marks[0]["Inizio"] if marks else None,
            "grid_start_beat_in_bar": marks[0]["Battito"] if marks else None,
            "meter": marks[0]["Metro"] if marks else None, "duration_sec": t.get("duration_sec"),
            "status": t.get("status"), "rb_track_id": attrs.get("TrackID")})
        if marks:
            arrays[idx] = {"tempo_map": encode([[m["Inizio"], m["Bpm"], int(m["Metro"].split("/")[0]), m["Battito"]]
                                                for m in marks], "<f8")}
    product = meta.get("product") or {}
    settings = meta.get("settings_from_user") or {}
    run = {"set": set_name, "program": f"{product.get('Name')} {product.get('Version')}",
           "company": product.get("Company"),
           "bpm_range": f"{settings.get('bpm_range')} (Analysis Setting, from the user; not stored in the XML)",
           "analysis_mode": f"{settings.get('analysis_mode')} (from the user; not stored in the XML)",
           "high_precision": f"{settings.get('high_precision')} (from the user; not stored in the XML)",
           "xml_precision": "AverageBpm and Bpm 2 decimals, Inizio 1 ms, TotalTime whole seconds",
           "xml": meta.get("source"), "xml_sha256": meta.get("source_sha256"),
           "xml_modified": meta.get("source_modified"), "playlist": meta.get("playlist"),
           "source": f"{jdir.name}/ (records of the collection XML)"}
    return run, rows, arrays


REKORDBOX = {
    "prefix": "rekordbox", "dir": P.REKORDBOX_DIR, "load": load_rekordbox,
    "groups": {"Track": "1F4E78", "BPM": "2E75B6", "Grid": "548235", "File": "595959"},
    "columns": [
        ("idx", "Track", "bench playlist position (1-1000)", "0"),
        ("file_name", "Track", "file name", None),
        ("bpm", "BPM", "AverageBpm rounded to 2 decimals", "0.00"),
        ("bpm_raw", "BPM", "AverageBpm as the XML reports it (2 decimals)", "0.00"),
        ("tempo_min", "BPM", "lowest Bpm among TEMPO markers", "0.00"),
        ("tempo_max", "BPM", "highest Bpm among TEMPO markers", "0.00"),
        ("n_tempo_markers", "BPM", "number of TEMPO markers (several markers can carry the same Bpm)", "0"),
        ("grid_start_sec", "Grid", "Inizio of the first TEMPO marker = first beat of the grid (XML: ms)", "0.000"),
        ("grid_start_beat_in_bar", "Grid", "Battito of the first marker: beat number in the bar (1 = downbeat)", "0"),
        ("meter", "Grid", "Metro of the first marker", None),
        ("duration_sec", "File", "TotalTime (XML: whole seconds)", "0"),
        ("status", "File", "ok = AverageBpm > 0", None),
        ("rb_track_id", "File", "TrackID in the Rekordbox collection", None),
        ("path", "File", "Location, decoded", None),
    ],
    "sqlite_only": [],
    "arrays": {"tempo_map": ("sec, bpm, beats, beat", "TEMPO markers [Inizio, Bpm, Metro numerator, Battito]; "
                                                       "beats of a segment are Inizio + k*60/Bpm until the next marker")},
    "notes": [("(amber rows)", "", "tempo_max differs from tempo_min: the Rekordbox grid changes tempo", "xlsx")],
}


# ---------------------------------------------------------------------------
# Beat This!
# ---------------------------------------------------------------------------
BT_FPS = 50
BT_RUN_KEYS = ("tool", "beat_this_version", "model", "postprocessing", "float16", "torch", "decoder",
               "playlist", "playlist_tracks", "last_started_at")


def load_beat_this(set_name: str, folder: Path):
    jdir = folder / f"json_{set_name}"
    meta = read_meta(jdir)
    rows, arrays = [], {}
    for idx, r in track_files(jdir):
        beats = np.asarray(r.get("beats_sec") or [], dtype="<f8")
        downs = np.asarray(r.get("downbeats_sec") or [], dtype="<f8")
        est = bpm_estimates(beats)
        hist = beats_per_bar(beats, downs)
        unfolded = est["bpm_median"]
        folded = fold_to_range(unfolded, BPM_LO, BPM_HI) if unfolded else None
        rows.append({
            "idx": idx, "file_name": file_name(r["path"]), "path": r["path"],
            "bpm": r2(folded), "bpm_raw": folded, "bpm_unfolded": unfolded,
            "n_beats": len(beats), "first_beat_sec": float(beats[0]) if len(beats) else None,
            "ibi_median_sec": est["ibi_median_sec"], "ibi_std_sec": est["ibi_std_sec"], "ibi_cv": est["ibi_cv"],
            "n_downbeats": len(downs), "first_downbeat_sec": float(downs[0]) if len(downs) else None,
            "beats_per_bar": int(max(hist, key=lambda k: hist[k])) if hist else None,
            "beats_per_bar_hist": "; ".join(f"{k}:{v}" for k, v in sorted(hist.items(), key=lambda kv: int(kv[0])))
            or None,
            "duration_sec": r.get("duration_sec"), "status": r.get("status")})
        arrays[idx] = {"beats_sec": encode(beats, "<f8"), "downbeats_sec": encode(downs, "<f8")}
    run = {"set": set_name, **{k: meta.get(k) for k in BT_RUN_KEYS}, "beat_fps": BT_FPS,
           "bpm_fold_range": "79-192 (bench MIK preset; applied by us, not by Beat This!)",
           "source": f"{jdir.name}/ (beat_this_bpm.py)"}
    return run, rows, arrays


BEAT_THIS = {
    "prefix": "beat_this", "dir": P.BEAT_THIS_DIR, "load": load_beat_this,
    "groups": {"Track": "1F4E78", "BPM": "2E75B6", "Beats": "548235", "Downbeats": "7030A0", "File": "595959"},
    "columns": [
        ("idx", "Track", "bench playlist position (1-1000)", "0"),
        ("file_name", "Track", "file name", None),
        ("bpm", "BPM", "bpm_raw rounded to 2 decimals", "0.00"),
        ("bpm_raw", "BPM", "60 / median inter-beat interval folded into 79-192 (Beat This! itself reports "
                           "only beats; quantized by its 50 fps grid)", "0.0000"),
        ("bpm_unfolded", "BPM", "60 / median inter-beat interval before folding", "0.0000"),
        ("n_beats", "Beats", "number of beats", "0"),
        ("first_beat_sec", "Beats", "first beat, s", "0.00"),
        ("ibi_median_sec", "Beats", "median inter-beat interval, s", "0.0000"),
        ("ibi_std_sec", "Beats", "standard deviation of inter-beat intervals, s", "0.0000"),
        ("ibi_cv", "Beats", "ibi_std_sec / mean interval", "0.0000"),
        ("n_downbeats", "Downbeats", "number of downbeats", "0"),
        ("first_downbeat_sec", "Downbeats", "first downbeat, s", "0.00"),
        ("beats_per_bar", "Downbeats", "most frequent number of beats between consecutive downbeats", "0"),
        ("beats_per_bar_hist", "Downbeats", "beats per bar : number of bars", None),
        ("duration_sec", "File", "decoded length (ffmpeg, 22050 Hz mono), 1 ms", "0.000"),
        ("status", "File", "ok = Beat This! returned beats", None),
        ("path", "File", "file Beat This! analysed", None),
    ],
    "sqlite_only": [],
    "arrays": {"beats_sec": ("sec", "beat times (Beat This! output, multiples of 1/50 s)"),
               "downbeats_sec": ("sec", "downbeat times (Beat This! output, multiples of 1/50 s)")},
    "notes": [("(amber bpm_unfolded)", "", "bpm_unfolded differs from bpm_raw: the tempo was folded into 79-192",
               "xlsx")],
}

PROGRAMS = {"sonara": SONARA, "mik": MIK, "rekordbox": REKORDBOX, "beat_this": BEAT_THIS}


# ---------------------------------------------------------------------------
# Writers
# ---------------------------------------------------------------------------
def write_sqlite(path: Path, spec: dict, run: dict, rows: list[dict], arrays: dict) -> None:
    tmp = path.with_suffix(".sqlite.tmp")
    if tmp.exists():
        tmp.unlink()
    con = sqlite3.connect(tmp)
    con.execute("CREATE TABLE run (key TEXT PRIMARY KEY, value TEXT)")
    con.executemany("INSERT INTO run VALUES (?, ?)",
                    [(k, v if isinstance(v, str) else json.dumps(v, ensure_ascii=False))
                     for k, v in {**run, "built_at": now_utc(), "built_by": "build_tables.py"}.items()])
    con.execute("CREATE TABLE fields (name TEXT, table_name TEXT, grp TEXT, meaning TEXT)")
    con.executemany("INSERT INTO fields VALUES (?, 'tracks', ?, ?)",
                    [(c, g, m) for c, g, m, _ in spec["columns"]] + list(spec["sqlite_only"]))
    con.executemany("INSERT INTO fields VALUES (?, 'arrays', 'Arrays', ?)",
                    [(n, f"{m}; unit: {u}") for n, (u, m) in spec["arrays"].items()])
    cols = [c for c, *_ in spec["columns"]] + [c for c, *_ in spec["sqlite_only"]]
    types = {c: ("INTEGER" if f == "0" else "REAL" if f else "TEXT") for c, _, _, f in spec["columns"]}
    col_defs = ", ".join(c + " " + types.get(c, "TEXT") for c in cols)
    con.execute(f"CREATE TABLE tracks ({col_defs}, PRIMARY KEY (idx))")
    con.executemany(f"INSERT INTO tracks VALUES ({','.join('?' * len(cols))})",
                    [[r.get(c) for c in cols] for r in rows])
    con.execute("CREATE TABLE arrays (idx INTEGER, name TEXT, dtype TEXT, shape TEXT, unit TEXT, "
                "data BLOB, PRIMARY KEY (idx, name))")
    con.executemany("INSERT INTO arrays VALUES (?, ?, ?, ?, ?, ?)",
                    [(i, n, dtype, shape, spec["arrays"][n][0], blob)
                     for i in sorted(arrays) for n, (dtype, shape, blob) in arrays[i].items()])
    con.commit()
    con.execute("VACUUM")
    con.close()
    os.replace(tmp, path)


def header_cell(ws, row: int, col: int, value, color: str) -> None:
    c = ws.cell(row, col, value)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = Border(left=WHITE_LINE, right=WHITE_LINE)


def write_xlsx(path: Path, set_name: str, spec: dict, run: dict, rows: list[dict]) -> None:
    columns, groups = spec["columns"], spec["groups"]
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Tracks"
    ws.sheet_properties.tabColor = "1F4E78" if set_name == "broken" else "7030A0"
    start = 0
    for j in range(1, len(columns) + 1):          # row 1: merged group headers
        if j == len(columns) or columns[j][1] != columns[start][1]:
            grp = columns[start][1]
            for k in range(start + 1, j + 1):
                header_cell(ws, 1, k, grp if k == start + 1 else None, groups[grp])
            if j > start + 1:
                ws.merge_cells(start_row=1, start_column=start + 1, end_row=1, end_column=j)
            start = j
    for j, (col, grp, _, _) in enumerate(columns, 1):  # row 2: field names
        header_cell(ws, 2, j, col, groups[grp])
    ws.row_dimensions[2].height = 45
    tempo_changes = spec is REKORDBOX
    n_col = [c for c, *_ in columns].index("n_tempo_markers") + 1 if tempo_changes else 0
    for i, r in enumerate(rows, 3):
        for j, (col, _, _, fmt) in enumerate(columns, 1):
            c = ws.cell(i, j, r.get(col))
            if fmt:
                c.number_format = fmt
            if tempo_changes and r.get("tempo_max") != r.get("tempo_min") and j <= n_col:
                c.fill = AMBER
            elif i % 2 == 0:
                c.fill = BAND
    last = len(rows) + 2
    letter = {col: get_column_letter(j) for j, (col, *_) in enumerate(columns, 1)}
    for j, (col, _, _, fmt) in enumerate(columns, 1):
        values = [r.get(col) for r in rows]
        if fmt and fmt != "0":
            dec = len(fmt.split(".")[1]) if "." in fmt else 0
            lens = [len(f"{v:.{dec}f}") for v in values if isinstance(v, (int, float))]
        else:
            lens = [len(str(v)) for v in values if v is not None]
        width = max([min(len(col), 14)] + lens) + 2
        ws.column_dimensions[get_column_letter(j)].width = min(width, {"file_name": 70, "path": 110}.get(col, 40))
    ws.freeze_panes = "C3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{last}"
    if "bpm_unfolded" in letter:
        u, b = letter["bpm_unfolded"], letter["bpm_raw"]
        ws.conditional_formatting.add(f"{u}3:{u}{last}", FormulaRule(formula=[f"ABS({u}3-{b}3)>0.0005"], fill=AMBER))
    for col in ("bpm_confidence", "grid_stability", "meter_grid_stability"):
        if col in letter:
            L = letter[col]
            ws.conditional_formatting.add(f"{L}3:{L}{last}", ColorScaleRule(
                start_type="num", start_value=0, start_color="F8696B", mid_type="num", mid_value=0.5,
                mid_color="FFEB84", end_type="num", end_value=1, end_color="63BE7B"))
    for col in ("rhythmic_regularity_label", "meter_rhythmic_regularity_label"):
        if col in letter:
            L = letter[col]
            ws.conditional_formatting.add(f"{L}3:{L}{last}", FormulaRule(formula=[f'{L}3="regular"'], fill=GREEN))
            ws.conditional_formatting.add(f"{L}3:{L}{last}", FormulaRule(formula=[f'{L}3="irregular"'], fill=AMBER))

    wf = wb.create_sheet("Fields")
    wf.sheet_properties.tabColor = "548235"
    for j, h in enumerate(("Field", "Group", "Meaning", "Where"), 1):
        header_cell(wf, 1, j, h, "1F4E78")
    entries = ([(c, g, m, "xlsx + sqlite") for c, g, m, _ in columns]
               + [(c, g, m, "sqlite only") for c, g, m in spec["sqlite_only"]]
               + [(n, "Arrays", f"{m}; unit: {u}", "sqlite, table arrays") for n, (u, m) in spec["arrays"].items()]
               + list(spec["notes"]))
    for i, e in enumerate(entries, 2):
        for j, v in enumerate(e, 1):
            c = wf.cell(i, j, v)
            if i % 2 == 1:
                c.fill = BAND
    for j, w in enumerate((38, 20, 95, 22), 1):
        wf.column_dimensions[get_column_letter(j)].width = w
    wf.freeze_panes = "A2"

    wr = wb.create_sheet("Run")
    wr.sheet_properties.tabColor = "595959"
    for j, h in enumerate(("Key", "Value"), 1):
        header_cell(wr, 1, j, h, "1F4E78")
    for i, (k, v) in enumerate({**run, "built_at": now_utc(), "built_by": "build_tables.py"}.items(), 2):
        wr.cell(i, 1, k).font = Font(bold=True)
        wr.cell(i, 2, v if v is None or isinstance(v, (str, int, float)) else json.dumps(v, ensure_ascii=False))
    wr.column_dimensions["A"].width = 24
    wr.column_dimensions["B"].width = 110
    wb.save(path)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--only", choices=tuple(PROGRAMS), action="append")
    ap.add_argument("--out-dir", help="write here instead of the program folders (for checks)")
    args = ap.parse_args()
    for key in args.only or PROGRAMS:
        spec = PROGRAMS[key]
        out = Path(args.out_dir) if args.out_dir else spec["dir"]
        out.mkdir(parents=True, exist_ok=True)
        for s in SETS:
            run, rows, arrays = spec["load"](s, spec["dir"])
            idxs = [r["idx"] for r in rows]
            if sorted(idxs) != list(range(1, len(idxs) + 1)):
                sys.exit(f"{key} {s}: track numbers are not 1..{len(idxs)}")
            db = out / f"{spec['prefix']}-{s}.sqlite"
            xlsx = out / f"{spec['prefix']}-{s}.xlsx"
            write_sqlite(db, spec, run, rows, arrays)
            write_xlsx(xlsx, s, spec, run, rows)
            print(f"{key} {s}: {len(rows)} tracks -> {db.name} {db.stat().st_size / 1e6:.1f} MB, "
                  f"{xlsx.name} {xlsx.stat().st_size / 1e6:.2f} MB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
